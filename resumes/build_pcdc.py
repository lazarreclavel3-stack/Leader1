"""Resume tailored to PCDC Director of Consulting Operations (ATS-friendly, 2 pages).

Usage: python3 resumes/build_pcdc.py out.docx
Reuses the layout helpers from build_executive_ats.py.
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

import build_executive_ats as ats
from build_executive_ats import ACCENT, MUTED, RIGHT_TAB, bottom_border, bullet, para, section, style_run

NAME = 'CLAVEL LAZARRE, MBA'
CONTACT = 'New York, NY  |  347-607-8164  |  lazarreclavel3@gmail.com  |  linkedin.com/in/clavel-lazarre'
HEADLINE = 'Consulting Operations  |  Finance  |  Program Delivery'

SUMMARY = (
    'Operations and finance leader with 20+ years of experience helping consulting practices, healthcare providers, and '
    'mission-driven organizations translate strategy into reliable daily execution. Brings a hands-on approach to '
    'connecting pricing, staffing capacity, client systems, and delivery requirements across fee-for-service work and '
    'grant-funded programs. Serves as a thought partner to leadership, making financial and operational information '
    'accessible, resolving cross-functional problems, and building practical systems that improve service quality and '
    'support sustainable growth.')

JOBS = [
    ('MAROON STRATEGIST, LLC', 'Brooklyn, NY', 'Fractional CFO & Operations Consultant', 'Mar 2024 – Present',
     'Fractional CFO and program leadership for nonprofit, public-sector, education, healthcare, and service '
     'organizations, coordinating financial planning and operating improvements across client engagements.',
     ['Manage hourly billing as part of daily consulting operations for multiple clients, including Atilem and Joseph '
      'Will Academy (JWA), alongside ongoing fractional CFO and project-based advisory work.',
      'Track client engagements in CRM systems, including Salesforce, and maintain Excel KPI tracking for grant-funded '
      'programs with daily, monthly, and quarterly compliance reviews and SkyNet engagement tracking.',
      'Support healthcare practice operations through accounting cleanup, tax-remediation coordination, staffing '
      'support, payroll and benefits processes, vendor oversight, and recurring financial reporting.',
      'Direct budgeting, forecasting, and cash management for client engagements, including multimillion-dollar project '
      'budgets with direct and indirect cost allocation and restricted cash-flow monitoring.',
      'Lead reconciliations and reporting remediation, coordinating clients, bookkeepers, accountants, and other '
      'stakeholders to resolve discrepancies and improve financial information.',
      'Build financial models, management dashboards, and standard operating procedures that help leadership assess '
      'performance, allocate resources, and improve delivery.',
      'Created and led the 18-month NYPA clean-energy workforce program from inception through completion, standing up '
      'design, staffing, budgeting, partnerships, and training; piloted the model before full rollout, and 92% of '
      'enrolled participants earned a certificate.',
      'Led proposal development and quality review for grant-funded and fee-for-service work; designed funder outcome '
      'reporting and formal partnership agreements with agencies, academic partners, and employers.']),
    ('OSBORNE ASSOCIATION', 'New York, NY', 'Vice President, Business Development & Social Ventures',
     'Jan 2023 – Oct 2025',
     'Led financial and operational oversight of a 140-person division, integrating staffing, finance, grant '
     'stewardship, partnerships, and program delivery across social ventures.',
     ['Managed 12 direct reports across finance, grant stewardship, operations, and frontline supervision; directed '
      'hiring, onboarding, training, and performance management.',
      'Ran weekly team meetings, individual check-ins, and recurring financial and operational reviews, using performance '
      'reports and staff feedback to identify problems, coordinate priorities, and guide decisions.',
      'Built pricing, monetization plans, and budgets for earned-revenue service lines, including landscaping and '
      'warehouse cleaning, as part of daily operations.',
      'Evaluated the feasibility of a proposed service expansion through a Rikers Island contract study, modeling '
      'delivery costs, staffing requirements, operational risks, and infrastructure needs.',
      'Held full P&L accountability for an $80M portfolio and ran staffing capacity for all 140 employees, managing '
      'hours, overtime, bonuses, and raises against budget.',
      'Managed multiple program budgets with monthly reconciliation and quarterly performance tracking, monitoring '
      'restricted and unrestricted cash flow.',
      'Led continuous improvement by modernizing payment processes, redesigning workflows, and introducing compliance '
      'meetings and staff self-assessments; helped move the division from a deficit exceeding $1M to a 13% operating '
      'surplus within one year.',
      'Led a nine-month transitional-employment program serving 300+ formerly incarcerated people annually, coordinating '
      'employer partnerships and tracking outcomes against funder benchmarks; partnered with the CEO and Board on growth.']),
    ('ALLUEL LOGISTICS', 'New York, NY', 'Vice President of Operations', 'Nov 2021 – Aug 2022',
     'Led daily operations across staffing, client systems, vendor management, billing, and public-sector contract '
     'delivery.',
     ['Led warehouse-management system selection and implementation, reducing data-processing time 60%; introduced CRM '
      'systems and trained staff to support adoption and consistent use.',
      'Directed quotation and invoicing processes across business and consumer engagements, including approximately '
      '$500K in monthly NYC DHS billing.',
      'Managed a $2M operating budget and negotiated vendor agreements that reduced supply-chain and logistics costs 35%.',
      'Coordinated delivery and compliance for NYCHA logistics contracts serving 100 developments, aligning operating '
      'activity with client requirements.',
      'Recruited and managed field, sales, and administrative staff; partnered with IT to improve HR workflows and data '
      'accuracy.']),
    ('LAZARRE CONSULTING GROUP', 'New York, NY', 'Founder & Principal Consultant', 'Dec 2016 – Present',
     'Project-based financial and business advisory practice, run alongside fractional CFO engagements, supporting '
     'private clients with financial analysis, operating improvements, and reporting.',
     ['Develop operational and performance-improvement strategies for organizations and brands including American Express '
      'OPEN, Ogilvy, and Chubb.',
      'Led financial remediation for a behavioral health practice, reconciling accounting records, identifying '
      'approximately $251K in unclaimed deductions, and preparing corrected financial statements.']),
    ('FITNESS COLLECTIVE INC.', 'Brooklyn, NY', 'Co-Founder', '2004 – 2016',
     'Built a service business from initial investment through expansion, establishing the staffing, pricing, client '
     'systems, and administrative infrastructure needed to operate across multiple locations.',
     ['Grew from a $5,000 investment to $1M in annual revenue within 24 months, then established a franchise model and '
      'launched three additional locations, expanding to $10M in annual sales.',
      'Created pricing for all services, connecting cost, staffing, and demand across personal training, group fitness, '
      'and other offerings.',
      'Scaled the workforce from four to 60 full-time employees across accounting, IT, facilities, personal training, and '
      'group fitness; ran all staff hours and scheduling, created roles, and participated directly in recruitment, '
      'supervision, and training.',
      'Selected and implemented Mindbody CRM for client tracking, booking, and targeted engagement across 4,000+ clients, '
      'using client data to shape services.',
      'Established a training academy supporting certification of 25 trainers; led website development, local '
      'partnerships, outreach, and vendor negotiations.']),
]

TECH = [
    ('SkyNet', 'Built and piloted an AI-supported workflow connecting research, evaluation, document preparation, and '
     'reporting, with evidence validation, cost controls, and performance tracking; now in full daily use.'),
    ('JWA AI Career Launch', 'Designed and submitted a Citi Foundation proposal for Joseph Will Academy connecting AI '
     'learning, workforce preparation, and employment support, with proposed measures for skills, employment, wages, '
     'and retention.'),
    ('Tools', 'Monday.com, QuickBooks, Mindbody CRM, Salesforce, Zoho CRM, Airtable, Excel, Google Workspace, and '
     'generative AI tools (ChatGPT/Codex, Claude, Gemini).'),
]

EDUCATION = [
    ('Master of Business Administration, General Business Management', 'West Texas A&M University', '2023'),
    ('Bachelor of Science, Small Business Administration', 'Northern Arizona University', ''),
]


def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.6)
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    doc.styles['Normal'].font.name, doc.styles['Normal'].font.size = ats.FONT, Pt(ats.BODY_PT)

    style_run(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(NAME), size=22, bold=True, color=ACCENT)
    style_run(para(doc, before=2, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(HEADLINE), size=11.5, bold=True)
    p = para(doc, before=2, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_run(p.add_run(CONTACT), size=10, color=MUTED)
    bottom_border(p, size=12)

    section(doc, 'Professional Summary')
    style_run(para(doc, after=1, line=1.05).add_run(SUMMARY))

    section(doc, 'Selected Professional Experience')
    for org, loc, title, dates, summary, bullets in JOBS:
        p = para(doc, before=5, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(org), bold=True, color=ACCENT)
        style_run(p.add_run(f'  |  {loc}'), color=MUTED)
        style_run(p.add_run('\t' + dates), bold=True)
        style_run(para(doc, after=1, keep_next=True).add_run(title), bold=True, italic=True)
        if summary:
            style_run(para(doc, after=1.5, keep_next=True).add_run(summary), size=10, italic=True, color=MUTED)
        for b in bullets:
            bullet(doc, b)

    section(doc, 'Technology & Program Development')
    for lead, text in TECH:
        bullet(doc, text, lead)

    section(doc, 'Education')
    for degree, school, year in EDUCATION:
        p = para(doc, after=1)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(degree), bold=True)
        style_run(p.add_run(f'  |  {school}'))
        if year:
            style_run(p.add_run('\t' + year))

    doc.core_properties.title = 'Clavel Lazarre, MBA — Resume'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build(sys.argv[1])
