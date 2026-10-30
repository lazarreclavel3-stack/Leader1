"""Resume and cover letter tailored to the CCM (Community Counseling & Mediation) CFO search.

Usage:
    python3 resumes/build_ccm.py resume out.docx
    python3 resumes/build_ccm.py letter out.docx
Reuses the layout helpers from build_executive_ats.py.
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

import build_executive_ats as ats
from build_executive_ats import ACCENT, MUTED, RIGHT_TAB, bottom_border, bullet, para, section, style_run

NAME = 'CLAVEL LAZARRE, MBA'
CONTACT = 'New York, NY  |  347-607-8164  |  lazarreclavel3@gmail.com  |  linkedin.com/in/clavel-lazarre'
HEADLINE = 'Nonprofit Finance & Operations Executive'
SUBHEAD = 'Financial Stewardship  |  Government Grants & Contracts  |  Behavioral Health & Human Services'

SUMMARY = (
    'Finance and operations executive with 20+ years of experience, including CFO-level leadership of nonprofit, '
    'healthcare, and public-sector finances. Held full P&L accountability for an $80M portfolio and financial oversight of '
    'a 140-person human-services division. Combines hands-on accounting oversight (fund and grant accounting, audit '
    'readiness, internal controls, and cash management) with strategic partnership to CEOs, boards, and program leaders. '
    'Brings deep experience with city and state human-services contracts, behavioral health advisory work, and earlier '
    'Medicare and Medicaid enrollment experience. Committed to strengthening organizations that expand equitable access to '
    'behavioral health care, supportive housing, and economic opportunity.')

JOBS = [
    ('MAROON STRATEGIST, LLC', 'Brooklyn, NY', 'Fractional Chief Financial Officer', 'Mar 2024 – Present',
     'Fractional CFO and operational leadership for nonprofit, healthcare, education, and service organizations, supporting '
     'financial planning, accounting improvements, compliance, and government-funded program delivery.',
     ['Direct budgeting, forecasting, treasury, cash-flow planning, and multi-entity financial analysis; evaluate staffing '
      'costs, operating expenses, and working-capital requirements to guide resource allocation.',
      'Direct audit readiness and external audit processes; develop internal controls and fund-accounting frameworks '
      'supporting GAAP, Uniform Guidance, and government funding requirements.',
      'Oversee accounts payable, accounts receivable, bank reconciliations, and monthly, quarterly, and annual financial '
      'statements for boards and funders; lead reporting remediation with bookkeepers and accounting advisers.',
      'Manage multimillion-dollar project budgets, allocating direct and indirect costs, monitoring restricted funds, and '
      'aligning expenditures with grant requirements and delivery milestones.',
      'Support behavioral health practice operations through accounting cleanup, financial controls, payroll and benefits '
      'processes, vendor oversight, and tax-remediation coordination.',
      'Develop financial models, management dashboards, and operating procedures that help leadership assess performance, '
      'identify risks, and implement corrective actions.',
      'Created and led the NYPA clean-energy workforce program from inception through delivery over 18 months, coordinating '
      'budgets, staffing, institutional partnerships, and training; 92% of enrolled participants completed training and '
      'earned a certificate.']),
    ('OSBORNE ASSOCIATION', 'New York, NY', 'Vice President, Business Development & Social Ventures', 'Jan 2023 – Oct 2025',
     'Led financial and operational oversight of a 140-person division within a nonprofit human-services organization, '
     'integrating financial stewardship, government-funded programs, staffing, and social-enterprise operations.',
     ['Held full P&L accountability for an $80M portfolio, aligning program budgets, staffing allocations, grant '
      'commitments, and operating requirements with organizational priorities.',
      'Helped move the division from a deficit exceeding $1M to a 13% operating surplus within one year through operational '
      'restructuring, cost controls, payment-process improvements, and revenue management.',
      'Managed multiple program budgets with monthly reconciliation and quarterly performance reviews; monitored restricted '
      'and unrestricted cash flow and identified financial risks and resource needs.',
      'Partnered with the CEO and Board, including the Finance and Audit Committees, on financial planning, risk, resource '
      'allocation, and program sustainability; negotiated government, corporate, and philanthropic partnerships.',
      'Supervised 12 direct reports across finance, grant stewardship, operations, and frontline management; directed '
      'hiring, onboarding, staff development, and performance management.',
      'Managed workforce capacity, staff hours, overtime, bonuses, and salary increases against budget; conducted recurring '
      'financial and operational reviews to coordinate decisions across teams.',
      'Oversaw IT systems, facilities, and vendor contracts supporting multi-site program operations.',
      'Developed service pricing and operating budgets for earned-revenue activities, including landscaping and warehouse '
      'cleaning; evaluated a proposed Rikers Island service expansion through cost, staffing, and infrastructure analysis.',
      'Oversaw a nine-month transitional-employment program serving 300+ formerly incarcerated individuals annually, '
      'coordinating employer partnerships and tracking performance against funder requirements.']),
    ('ALLUEL LOGISTICS', 'New York, NY', 'Vice President of Operations', 'Nov 2021 – Aug 2022',
     'Directed daily operations across budgeting, public-sector contracts, billing, staffing, technology implementation, '
     'and vendor management.',
     ['Managed a $2M operating budget and negotiated vendor agreements that reduced logistics and supply-chain costs 35%.',
      'Directed financial controls, quotation and invoicing processes, and contract billing, including approximately $500K '
      'in monthly NYC Department of Homeless Services billing.',
      'Oversaw NYCHA logistics contracts supporting 100 developments, aligning daily operations with contractual obligations '
      'and client requirements.',
      'Led warehouse-management system selection and implementation, reducing data-processing time 60%; introduced CRM '
      'systems and trained staff to improve adoption and data accuracy.']),
    ('LAZARRE CONSULTING GROUP', 'New York, NY', 'Founder & Fractional CFO', 'Dec 2016 – Present',
     'Project-based financial and business advisory practice supporting financial analysis, cash management, reporting, '
     'and organizational improvement.',
     ['Provide cash-flow planning, financial analysis, and operational guidance; translate financial findings into '
      'actionable priorities for client leadership and accounting advisers.',
      'Led financial remediation for a behavioral health practice, reconciling accounting records, identifying '
      'approximately $251K in unclaimed deductions, and preparing corrected financial statements.',
      'Supported capital planning, financing strategy, and banking relationships for a $200M national banking development '
      'initiative; developed performance-improvement strategies for organizations including American Express OPEN, Ogilvy, '
      'and Chubb.']),
    ('FITNESS COLLECTIVE INC.', 'Brooklyn, NY', 'Co-Founder & Chief Executive Officer', '2004 – 2016',
     'Built and led a multi-location health and wellness business, establishing its financial, staffing, technology, and '
     'facilities infrastructure.',
     ['Grew a $5,000 investment into $1M in annual revenue within 24 months, then established a franchise model and expanded '
      'to $10M in annual sales.',
      'Scaled the workforce from four to 60 full-time employees across accounting, IT, facilities, training, and operations; '
      'oversaw recruitment, supervision, scheduling, payroll, and benefits.',
      'Developed service pricing based on cost, staffing, and demand; managed facilities, technology, insurance, and vendor '
      'relationships across locations.']),
]

HEALTH = [
    'Advised Medicare and Medicaid beneficiaries on Aetna plan benefits, coverage options, and enrollment.',
    'Worked across health, commercial, life, and accident insurance with hospitals and agents; previously held Life, '
    'Accident & Health insurance licenses.',
]

SYSTEMS = ('QuickBooks  |  Excel financial models and KPI reporting  |  Monday.com  |  Salesforce  |  Zoho CRM  |  '
           'Airtable  |  Mindbody  |  Google Workspace')
SKYNET = ('Developed and piloted SkyNet, an AI-supported workflow now in daily use for research, evaluation, document '
          'preparation, and engagement tracking, with evidence validation, cost controls, and performance monitoring.')

EDUCATION = [
    ('Master of Business Administration, General Business Management', 'West Texas A&M University', '2023'),
    ('Bachelor of Science, Small Business Administration', 'Northern Arizona University', ''),
]

LETTER = [
    'Dear Mr. Spivey and the CCM Search Committee,',
    'I am writing to apply for the position of Chief Financial Officer at Community Counseling & Mediation. For more than '
    'four decades, CCM has met Brooklyn families where they are, combining mental health, substance use, supportive '
    'housing, and youth services in communities that are too often underserved. That integrated model depends on a finance '
    'function that keeps more than 75 programs funded, compliant, and able to grow. I would welcome the opportunity to lead '
    'that work.',
    'I bring 20+ years of finance and operations leadership, including CFO-level responsibility for nonprofit, healthcare, '
    'and public-sector organizations. At the Osborne Association, I held full P&L accountability for an $80M portfolio and '
    'financial oversight of a 140-person human-services division. Working with the CEO and the Board’s Finance and Audit '
    'Committees, I helped move the division from a deficit exceeding $1M to a 13% operating surplus within one year '
    'through restructuring, cost controls, payment-process improvements, and disciplined revenue management. I managed '
    'multiple program budgets with monthly reconciliation and quarterly reviews, monitored restricted and unrestricted '
    'cash flow, led 12 direct reports, and oversaw IT systems, facilities, and vendor contracts across sites.',
    'Government contract compliance has been central to my work. I directed billing and controls for approximately $500K '
    'in monthly NYC Department of Homeless Services billing and NYCHA contracts serving 100 developments, and I built the '
    'financial framework for an 18-month New York Power Authority workforce program, allocating direct and indirect costs '
    'and aligning restricted spending with grant requirements. As a fractional CFO, I direct audit readiness, internal '
    'controls, fund-accounting frameworks, AP/AR, bank reconciliations, and monthly, quarterly, and annual financial '
    'statements for boards and funders.',
    'Behavioral health is also familiar ground. I currently support a behavioral health practice’s financial operations, '
    'and I previously led a remediation that reconciled its records, identified approximately $251K in unclaimed '
    'deductions, and produced corrected financial statements. Earlier in my career, I advised Medicare and Medicaid '
    'beneficiaries on plan benefits and enrollment, which gives me an appreciation for how reimbursement shapes the people '
    'and programs CCM serves.',
    'I would welcome the opportunity to discuss how I can strengthen CCM’s financial stewardship and support its next '
    'stage of growth. Thank you for your time and consideration.',
]


def page_setup(doc, lr, tb, size):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(lr)
    sec.top_margin = sec.bottom_margin = Inches(tb)
    doc.styles['Normal'].font.name, doc.styles['Normal'].font.size = ats.FONT, Pt(size)


def header(doc, with_headline):
    style_run(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(NAME), size=22, bold=True, color=ACCENT)
    if with_headline:
        style_run(para(doc, before=2, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(HEADLINE), size=11.5, bold=True)
        style_run(para(doc, before=1, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(SUBHEAD), size=10, color=ACCENT)
    p = para(doc, before=3, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_run(p.add_run(CONTACT), size=10, color=MUTED)
    bottom_border(p, size=12)


def build_resume(out):
    doc = Document()
    page_setup(doc, 0.6, 0.5, ats.BODY_PT)
    header(doc, True)

    section(doc, 'Professional Summary')
    style_run(para(doc, after=1, line=1.05).add_run(SUMMARY))

    section(doc, 'Selected Professional Experience')
    for org, loc, title, dates, summary, bullets in JOBS:
        p = para(doc, before=5, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(org), bold=True, color=ACCENT)
        style_run(p.add_run(f'  |  {loc}'), color=MUTED)
        style_run(p.add_run('\t' + dates), bold=True)
        style_run(para(doc, after=1, keep_next=True).add_run(title), bold=True, italic=True)
        style_run(para(doc, after=1.5, keep_next=True).add_run(summary), size=10, italic=True, color=MUTED)
        for b in bullets:
            bullet(doc, b)

    section(doc, 'Additional Healthcare & Insurance Experience')
    p = para(doc, before=2, keep_next=True)
    style_run(p.add_run('PFP FINANCIAL / PFP INSURANCE'), bold=True, color=ACCENT)
    style_run(p.add_run('  |  Management Consulting, Business Development & Insurance'), color=MUTED)
    for b in HEALTH:
        bullet(doc, b)

    section(doc, 'Financial Systems & Technology')
    style_run(para(doc, after=2).add_run(SYSTEMS))
    style_run(para(doc, after=1, line=1.03).add_run(SKYNET))

    section(doc, 'Education')
    for degree, school, year in EDUCATION:
        p = para(doc, after=1)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(degree), bold=True)
        style_run(p.add_run(f'  |  {school}'))
        if year:
            style_run(p.add_run('\t' + year))

    doc.core_properties.title = 'Clavel Lazarre, MBA — Resume (CFO)'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)


def build_letter(out):
    doc = Document()
    page_setup(doc, 0.9, 0.7, 11)
    header(doc, False)
    for i, text in enumerate(LETTER):
        style_run(para(doc, before=12 if i == 0 else 0, after=8, line=1.1).add_run(text), size=11)
    style_run(para(doc, before=2, after=2).add_run('Sincerely,'), size=11)
    style_run(para(doc, before=12).add_run('Clavel Lazarre, MBA'), size=11, bold=True)
    doc.core_properties.title = 'Clavel Lazarre, MBA — Cover Letter, CCM CFO'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)


if __name__ == '__main__':
    {'resume': build_resume, 'letter': build_letter}[sys.argv[1]](sys.argv[2])
    print(sys.argv[2])
