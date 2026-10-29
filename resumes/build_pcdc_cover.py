"""Cover letter for PCDC Director of Consulting Operations, matching the resume style.

Usage: python3 resumes/build_pcdc_cover.py out.docx
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

import build_executive_ats as ats
from build_executive_ats import ACCENT, MUTED, bottom_border, para, style_run

CONTACT = 'New York, NY  |  347-607-8164  |  lazarreclavel3@gmail.com  |  linkedin.com/in/clavel-lazarre'

PARAGRAPHS = [
    'Dear PCDC Hiring Team,',
    'Access to quality healthcare should not depend on a person’s ZIP code, income, or ability to navigate a complicated '
    'system. Too many people in under-resourced communities are denied not only care, but the dignity of being heard, '
    'respected, and supported. PCDC’s work as a CDFI, financing and advising the safety-net providers who serve those '
    'communities, is why I am applying for the Director of Consulting Operations role. It offers an opportunity to put my '
    'operational and financial experience directly behind that conviction.',
    'Across business leadership, workforce development, and healthcare advisory work, I have focused on expanding access '
    'and helping organizations deliver on their commitments. At Osborne, that meant strengthening businesses and '
    'employment pathways for people returning from incarceration. Through the NYPA clean-energy workforce program, it '
    'meant helping participants gain skills and access new career opportunities. In my current healthcare advisory work, '
    'it means helping a practice address the financial, staffing, and administrative challenges that affect its ability '
    'to serve patients. Each experience reinforced the same lesson: access depends on organizations having the resources, '
    'systems, and people to sustain it.',
    'I have done this operational work directly. At Osborne, I built pricing, monetization plans, and budgets for '
    'earned-revenue service lines, managed staffing hours and compensation for a 140-person division, and led 12 direct '
    'reports through weekly operating and financial reviews. In my consulting practice, I manage hourly billing and track '
    'client engagements in CRM systems. I have led system rollouts from selection through launch, including a '
    'warehouse-management implementation that cut processing time 60%. I also stood up the NYPA program from the ground '
    'up, piloting the model before full rollout and reaching a 92% certificate-completion rate.',
    'What appeals to me most about this role is the connection between purpose and execution. A consulting practice that '
    'serves both fee-for-service clients and grant-funded programs needs thoughtful pricing, realistic capacity planning, '
    'reliable engagement tracking, and disciplined financial management to do its best work. After years of leading at '
    'the executive level, I am seeking a role where I own that operating engine day to day: bringing clarity to complex '
    'situations, supporting colleagues, and building practical processes that make their work easier and their impact '
    'more durable.',
    'I would welcome the opportunity to discuss how I can help PCDC’s consulting practice operate with clarity and grow '
    'sustainably. Thank you for your time and consideration.',
]


def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    doc.styles['Normal'].font.name, doc.styles['Normal'].font.size = ats.FONT, Pt(11)

    style_run(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER).add_run('CLAVEL LAZARRE, MBA'), size=22, bold=True, color=ACCENT)
    p = para(doc, before=3, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_run(p.add_run(CONTACT), size=10, color=MUTED)
    bottom_border(p, size=12)

    for i, text in enumerate(PARAGRAPHS):
        p = para(doc, before=14 if i == 0 else 0, after=9, line=1.12)
        style_run(p.add_run(text), size=11)

    style_run(para(doc, before=4, after=2).add_run('Sincerely,'), size=11)
    style_run(para(doc, before=14).add_run('Clavel Lazarre, MBA'), size=11, bold=True)

    doc.core_properties.title = 'Clavel Lazarre, MBA — Cover Letter, PCDC'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build(sys.argv[1])
