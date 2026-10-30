"""Resume and cover letter tailored to the EngageWell (CHIC / EngageWell IPA) Executive Director search.

Usage:
    python3 resumes/build_engagewell.py resume out.docx
    python3 resumes/build_engagewell.py letter out.docx
Reuses the layout helpers from build_executive_ats.py and build_ccm.py.
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

import build_executive_ats as ats
from build_executive_ats import ACCENT, MUTED, RIGHT_TAB, bottom_border, bullet, para, section, style_run
from build_ccm import CONTACT, NAME, page_setup

HEADLINE = 'Nonprofit Executive  |  Healthcare Finance & Revenue Growth  |  Health Equity'

SUMMARY = (
    'Mission-driven finance and operations executive with 20+ years of experience growing revenue, managing healthcare '
    'and nonprofit finances, and building organizations that serve marginalized communities. Brings hands-on experience '
    'with IPA-related finances and Medicaid managed care billing, substance use and reentry programs, Form 990 and audit '
    'oversight, and government contracts. An entrepreneurial builder who partners with boards, runs day-to-day finance and '
    'administration in lean organizations, and translates relationships into sustainable funding.')

JOBS = [
    ('MAROON STRATEGIST, LLC', 'Brooklyn, NY', 'Fractional Chief Financial Officer', 'Mar 2024 – Present',
     'Fractional CFO and operational leadership for healthcare, nonprofit, education, and service organizations.',
     ['Serve as fractional CFO for a physician practice participating in an IPA, managing IPA-related finances and billing '
      'to Medicaid managed care and other insurance plans.',
      'Oversee payroll, accounts payable, vendor payments, invoicing, bank reconciliations, and monthly, quarterly, and '
      'annual financial statements for boards and funders.',
      'Manage hourly billing and client engagement tracking in CRM systems, including Salesforce, for multiple consulting '
      'clients alongside fractional CFO work.',
      'Direct budgeting, forecasting, cash-flow planning, and internal controls; develop fund-accounting frameworks '
      'supporting GAAP, Uniform Guidance, and funder requirements.',
      'Lead reconciliations and financial-reporting remediation, coordinating clients, bookkeepers, and accounting '
      'advisers to resolve discrepancies and improve the reliability of financial records.',
      'Lead annual audit coordination and readiness testing for workforce development, healthcare programs, and smaller '
      '501(c)(3) organizations.',
      'Support behavioral health practice operations through accounting cleanup, financial controls, payroll and benefits '
      'processes, vendor oversight, and tax-remediation coordination.',
      'As an IRS-registered tax preparer, compile Form 990 returns for multiple nonprofits and prepare business and '
      'individual returns, including dozens of state and local filings.',
      'Created and led the NYPA clean-energy workforce program from inception through delivery over 18 months, building a '
      'coalition of public agencies, elected officials, Stanford Innovation Lab, KPMG, NYSID, and employers; piloted the '
      'model before full rollout, and 92% of enrolled participants completed training and earned a certificate.',
      'Lead proposal development and partnership agreements for grant-funded and fee-for-service work; track program KPIs, '
      'restricted funding, and funder compliance through monthly and quarterly reviews.']),
    ('OSBORNE ASSOCIATION', 'New York, NY', 'Vice President, Business Development & Social Ventures', 'Jan 2023 – Oct 2025',
     'Led business development and financial oversight of a 140-person division in a human-services organization serving '
     'people affected by the criminal legal system.',
     ['Led fundraising for the division: identified and prioritized prospects, developed proposals, and personally led '
      'solicitation of the highest-value donors, foundations, and corporate partners, raising $400K+ from funders and '
      'partners including REDF and JC Flowers.',
      'Partnered with the CEO, Board, and Fund Development Committee on donor cultivation and stewardship, deploying board '
      'members in cultivation and co-leading annual fundraising events that exceeded revenue targets.',
      'Structured and negotiated $10M+ in new government, institutional, and corporate contracts, turning partnerships '
      'into funded programs and employment pathways.',
      'Built pricing, monetization plans, and budgets for earned-revenue social ventures, including landscaping and '
      'warehouse cleaning, diversifying the division’s revenue.',
      'Helped move the division from a deficit exceeding $1M to a 13% operating surplus within one year through '
      'restructuring, cost controls, payment-process improvements, and revenue management.',
      'Led continuous improvement and data-informed decision-making, modernizing payment processes, redesigning '
      'workflows, and introducing compliance meetings, staff self-assessments, and weekly operating reviews.',
      'Evaluated a proposed Rikers Island service expansion, modeling delivery costs, staffing, operational risks, and '
      'infrastructure needs, and developed a roadmap for potentially doubling a program within 12 months.',
      'Held full P&L accountability for an $80M portfolio; partnered with the CEO and Board, including the Finance and '
      'Audit Committees, on planning, risk, and sustainability.',
      'Partnered with the Director of El Rio, a substance use treatment program for people returning from incarceration, '
      'on budgeting, hiring, and participant placements through 2024.',
      'Oversaw a nine-month transitional-employment program serving 300+ formerly incarcerated people annually, '
      'coordinating employer partnerships and tracking outcomes against funder requirements.',
      'Managed multiple program budgets with monthly reconciliation and quarterly performance reviews, monitoring '
      'restricted and unrestricted cash flow.',
      'Supervised 12 direct reports; managed staff hours, overtime, bonuses, and raises against budget; oversaw IT systems, '
      'facilities, and vendor contracts.']),
    ('ALLUEL LOGISTICS', 'New York, NY', 'Vice President of Operations', 'Nov 2021 – Aug 2022', None,
     ['Directed contract billing and financial controls, including approximately $500K in monthly NYC Department of '
      'Homeless Services billing and NYCHA contracts serving 100 developments.',
      'Managed a $2M operating budget and negotiated vendor agreements that reduced costs 35%; led a warehouse-management '
      'system rollout that cut processing time 60%.',
      'Recruited and managed field, sales, and administrative staff; partnered with IT to improve HR workflows and data '
      'accuracy.']),
    ('LAZARRE CONSULTING GROUP', 'New York, NY', 'Founder & Fractional CFO', 'Dec 2016 – Present', None,
     ['Led financial remediation for a behavioral health practice, reconciling records, identifying approximately $251K in '
      'unclaimed deductions, and preparing corrected financial statements.',
      'Supported capital planning and financing strategy for a $200M national banking development initiative; advised '
      'organizations including American Express OPEN, Ogilvy, and Chubb.']),
    ('FITNESS COLLECTIVE INC.', 'Brooklyn, NY', 'Co-Founder & Chief Executive Officer', '2004 – 2016', None,
     ['Grew a $5,000 investment into $1M in annual revenue within 24 months, then franchised and expanded to $10M in annual '
      'sales.',
      'Scaled the workforce from four to 60 full-time employees; built HR, payroll, and benefits systems and established '
      'service pricing, staffing, and operating processes.',
      'Built neighborhood partnerships, outreach, and community events to establish a presence in each location and grow a '
      'client base of 4,000+.',
      'Established a training academy that certified 25 trainers, expanding workforce capacity and service offerings.']),
]

HEALTH = [
    'Advised Medicare and Medicaid beneficiaries on Aetna plan benefits and enrollment, worked with hospitals and '
    'health-insurance agents, and sold commercial insurance to organizations; previously held Life, Accident & Health '
    'insurance licenses.',
]

SYSTEMS = ('QuickBooks  |  Excel financial models and KPI reporting  |  Monday.com  |  Salesforce  |  Zoho CRM  |  '
           'Airtable  |  Google Workspace')
SKYNET = ('Developed and piloted SkyNet, an AI-supported workflow now in daily use for research, evaluation, document '
          'preparation, and engagement tracking.')

EDUCATION = [
    ('Master of Business Administration, General Business Management', 'West Texas A&M University', '2023'),
    ('Bachelor of Science, Small Business Administration', 'Northern Arizona University', ''),
]

LETTER = [
    'Dear EngageWell Search Committee,',
    'I am writing to apply for the position of Executive Director of EngageWell Community Health & Innovation '
    'Collaborative and EngageWell IPA. EngageWell’s work, integrating community-based organizations into how healthcare is '
    'delivered and financed, with a commitment to harm reduction and to the dignity of the people it serves, reflects '
    'what I have spent my career building toward. I would welcome the opportunity to lead its next chapter of growth.',
    'Revenue diversification has been central to my work. At the Osborne Association, I structured and negotiated more '
    'than $10M in new government, institutional, and corporate contracts and raised over $400K from private donors, '
    'foundations, and corporate partners, including REDF and JC Flowers, by identifying prospects, developing proposals, '
    'personally leading solicitations, and stewarding those relationships with the CEO and Board. I also built pricing, monetization plans, and '
    'budgets for earned-revenue social ventures, and I helped move a 140-person division from a deficit exceeding $1M to a '
    '13% operating surplus within one year while partnering with the CEO and the Board’s Finance and Audit Committees.',
    'I also bring hands-on healthcare financing experience. As a fractional CFO, I manage the finances of a physician '
    'practice that participates in an IPA, including its IPA-related reporting and billing to Medicaid managed care and '
    'other insurance plans. Earlier in my career, I advised Medicare and Medicaid beneficiaries on plan benefits and '
    'enrollment. At Osborne, I partnered with the Director of El Rio, a substance use treatment program for people '
    'returning from incarceration, on budgeting, hiring, and participant placements. That work taught me that people '
    'engage with care when they are met where they are, treated with dignity, and trusted to make their own choices. '
    'Nonjudgmental, low-barrier access is not only a value; it is how programs earn the trust that keeps people '
    'connected to services.',
    'A small organization needs a leader who can set strategy and also run the numbers. I currently oversee payroll, '
    'accounts payable, invoicing, reconciliations, and financial statements for boards and funders; lead annual audit '
    'coordination; and, as an IRS-registered tax preparer, compile Form 990 returns for multiple nonprofits. I have also '
    'built organizations from the ground up, from a fitness business I grew from a $5,000 investment to $10M in annual '
    'sales to an 18-month NYPA workforce program in which 92% of participants earned a certificate.',
    'I would welcome the opportunity to discuss how I can strengthen EngageWell’s financial sustainability, grow its '
    'network, and advance its mission. Thank you for your time and consideration.',
]


def header(doc, with_headline):
    style_run(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(NAME), size=22, bold=True, color=ACCENT)
    if with_headline:
        style_run(para(doc, before=2, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(HEADLINE), size=11.5, bold=True)
    p = para(doc, before=3, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_run(p.add_run(CONTACT), size=10, color=MUTED)
    bottom_border(p, size=12)


def build_resume(out):
    doc = Document()
    page_setup(doc, 0.6, 0.5, ats.BODY_PT)
    header(doc, True)

    section(doc, 'Professional Summary')
    style_run(para(doc, after=1, line=1.05).add_run(SUMMARY))

    section(doc, 'Selected Professional Experience')
    for org, loc, title, dates, summary, bullets in JOBS:
        p = para(doc, before=5, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(org), bold=True, color=ACCENT)
        style_run(p.add_run(f'  |  {loc}'), color=MUTED)
        style_run(p.add_run('\t' + dates), bold=True)
        style_run(para(doc, after=1, keep_next=True).add_run(title), bold=True, italic=True)
        if summary:
            style_run(para(doc, after=1.5, keep_next=True).add_run(summary), size=10, italic=True, color=MUTED)
        for b in bullets:
            bullet(doc, b)

    section(doc, 'Additional Healthcare & Insurance Experience')
    p = para(doc, before=2, keep_next=True)
    style_run(p.add_run('PFP FINANCIAL / PFP INSURANCE'), bold=True, color=ACCENT)
    style_run(p.add_run('  |  Management Consulting, Business Development & Insurance'), color=MUTED)
    for b in HEALTH:
        bullet(doc, b)

    section(doc, 'Financial Systems & Technology')
    style_run(para(doc, after=2).add_run(SYSTEMS))
    style_run(para(doc, after=1, line=1.03).add_run(SKYNET))

    section(doc, 'Education')
    for degree, school, year in EDUCATION:
        p = para(doc, after=1)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(degree), bold=True)
        style_run(p.add_run(f'  |  {school}'))
        if year:
            style_run(p.add_run('\t' + year))

    doc.core_properties.title = 'Clavel Lazarre, MBA — Resume (Executive Director)'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)


def build_letter(out):
    doc = Document()
    page_setup(doc, 0.9, 0.7, 11)
    header(doc, False)
    for i, text in enumerate(LETTER):
        style_run(para(doc, before=12 if i == 0 else 0, after=8, line=1.1).add_run(text), size=11)
    style_run(para(doc, before=2, after=2).add_run('Sincerely,'), size=11)
    style_run(para(doc, before=12).add_run('Clavel Lazarre, MBA'), size=11, bold=True)
    doc.core_properties.title = 'Clavel Lazarre, MBA — Cover Letter, EngageWell Executive Director'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)


if __name__ == '__main__':
    {'resume': build_resume, 'letter': build_letter}[sys.argv[1]](sys.argv[2])
    print(sys.argv[2])
