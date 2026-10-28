"""Two-page executive resume, ATS-friendly best-practice layout.

Single column, standard headings, no tables/text boxes/graphics, dates on
right-aligned tab stops, Calibri 10.5pt body, one restrained accent color.

Usage: python3 resumes/build_executive_ats.py out.docx
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)   # deep navy
TEXT = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
FONT = 'Calibri'
BODY_PT = 10.5
RIGHT_TAB = Inches(7.2)

NAME = 'CLAVEL LAZARRE, MBA'
CONTACT = 'New York, NY  |  347-607-8164  |  lazarreclavel3@gmail.com  |  linkedin.com/in/clavel-lazarre'
HEADLINE = 'Executive Leadership  |  Organizational Growth  |  Strategy & Operations'

SUMMARY = [
    'Operations and finance executive with 20+ years of experience building businesses, leading mission-driven '
    'programs, and creating the systems organizations need to grow. Held full P&L accountability for an $80M '
    'portfolio and financial and operational oversight of a 140-person division. Co-founded a fitness business that '
    'reached $1M in annual revenue within 24 months, then expanded through a franchise model to $10M in annual sales.',
    'Brings a founder’s understanding of growth and an operator’s discipline in execution. Partners with CEOs, boards, '
    'public agencies, and community organizations to evaluate opportunities, align resources, and translate strategy '
    'into budgets, partnerships, programs, and operating processes.',
]

RESULTS = [
    ('Built and expanded a business', 'Grew Fitness Collective from a $5,000 investment to $10M in annual sales, '
     'establishing a franchise model and launching three additional locations.'),
    ('Led at substantial scale', 'Oversaw a 140-person division and an $80M P&L at Osborne Association, connecting '
     'financial accountability, workforce development, partnerships, and service delivery.'),
    ('Created mission-driven programs', 'Led a nine-month transitional-employment program serving 300+ formerly '
     'incarcerated people annually; created and delivered an 18-month NYPA clean-energy workforce project.'),
    ('Improved operating performance', 'Cut vendor costs 35% and processing time 60% at Alluel Logistics through '
     'negotiation and systems improvements.'),
]

EXPERTISE = ['Organizational Growth', 'General Management', 'P&L Leadership', 'Budgeting & Resource Allocation',
             'CEO & Board Partnership', 'People Operations', 'Program Development', 'Coalition Building',
             'Fundraising & Grant Stewardship', 'Risk Management', 'Performance Reporting & KPIs',
             'Technology & Process Improvement', 'Workforce Development', 'Social Enterprise', 'Practical AI Implementation']

JOBS = [
    ('MAROON STRATEGIST, LLC', 'Brooklyn, NY', 'Chief Financial Strategy Officer', 'Mar 2024 – Present',
     'Fractional CFO and program leadership for nonprofit, public-sector, education, healthcare, and service businesses.',
     ['Created and led the NYPA clean-energy workforce training project from inception through final delivery over 18 '
      'months, owning proposal development, program design, budgeting, partner coordination, and participant training.',
      'Built a workforce coalition spanning city and state agencies, elected officials, academic partners, Stanford '
      'Innovation Lab, KPMG, and the New York Power Authority.',
      'Managed a multimillion-dollar project budget: allocated direct and indirect costs, managed restricted cash flow, '
      'and aligned spending with grant requirements and delivery milestones.',
      'Led fundraising and partnership-agreement development, tying funding commitments to partner responsibilities and '
      'implementation needs.',
      'Direct budgeting, forecasting, cash management, and multi-entity financial planning; advise leaders on financing, '
      'working capital, and growth priorities.',
      'Build financial models, KPI dashboards, and SOPs; lead financial cleanup and reporting remediation with clients '
      'and their accountants.']),
    ('OSBORNE ASSOCIATION', 'New York, NY', 'Vice President, Business Development & Social Ventures',
     'Jan 2023 – Oct 2025',
     'Led a 140-person division of a human-services organization supporting people affected by the criminal legal system.',
     ['Held full P&L accountability for an $80M portfolio, aligning budgets, staffing, partnerships, and program priorities.',
      'Partnered with the CEO and Board on long-range planning, risk management, resource allocation, and growth.',
      'Led a nine-month transitional-employment program serving 300+ formerly incarcerated people annually, coordinating '
      'employer partnerships and tracking outcomes against funder benchmarks.',
      'Negotiated government, corporate, and philanthropic partnerships and launched earned-income social ventures.',
      'Authored a growth and feasibility white paper on a potential Rikers Island service contract, integrating financial '
      'modeling, workforce needs, mission alignment, employee well-being, and operational risk.',
      'Developed a roadmap to potentially double JMS within 12 months, sizing investments in management capacity, '
      'recruitment, training, contract oversight, payroll, compliance, and technology.',
      'Oversaw departmental talent acquisition, onboarding, staff development, and performance management.']),
    ('ALLUEL LOGISTICS', 'New York, NY', 'Vice President of Operations', 'Nov 2021 – Aug 2022', None,
     ['Managed a $2M operating budget; renegotiated vendor agreements to cut supply-chain and logistics costs 35%.',
      'Implemented warehouse-management technology that reduced data-processing time 60%; introduced CRM systems and '
      'trained staff.',
      'Oversaw NYCHA logistics contracts serving 100 developments and supervised ~$500K in monthly NYC DHS billing.',
      'Recruited and managed field, sales, and administrative staff; improved HR workflows and data accuracy.']),
    ('LAZARRE CONSULTING GROUP', 'New York, NY', 'Chief Executive Officer', 'Dec 2016 – Present', None,
     ['Served as financial architect for a $200M national banking development initiative, supporting capital planning '
      'and financial structuring.',
      'Developed operating and performance-improvement strategies for organizations and brands including American '
      'Express OPEN, Ogilvy, and Chubb.',
      'Led financial remediation for a psychological services practice, identifying ~$251K in unclaimed deductions and '
      'preparing corrected financial statements.']),
    ('FITNESS COLLECTIVE INC.', 'Brooklyn, NY', 'CEO & Co-Founder', '2004 – 2016', None,
     ['Scaled from a $5,000 investment to $1M in annual revenue within 24 months, then franchised and launched three '
      'additional locations, growing to $10M in annual sales.',
      'Built HR, payroll, and benefits infrastructure for 50+ employees, overseeing recruitment, training, performance '
      'management, and ~$65K in monthly payroll.',
      'Implemented Mindbody CRM for client tracking, booking, and targeted engagement across 4,000+ clients; founded a '
      'training academy that certified 25 trainers.',
      'Developed local partnerships, cross-marketing, and community events; negotiated vendor and utility agreements.']),
]

INNOVATION = [
    ('SkyNet', 'Built an AI-supported workflow connecting opportunity research, evaluation, document preparation, and '
     'reporting, with evidence validation, cost controls, and performance tracking.'),
    ('JWA AI Career Launch', 'Designed and submitted a Citi Foundation proposal linking AI learning, workforce '
     'preparation, and employment support, with measures for skills, employment, wages, and retention.'),
]

EDUCATION = [
    ('Master of Business Administration, General Business Management', 'West Texas A&M University', '2023'),
    ('Bachelor of Science, Small Business Administration', 'Northern Arizona University', ''),
]


def style_run(r, size=BODY_PT, bold=False, italic=False, color=TEXT, caps=False):
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.all_caps = caps
    return r


def para(doc, before=0, after=0, line=1.0, align=None, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    return p


def bottom_border(p, color='1F3A5F', size=8):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', str(size)), ('w:space', '1'), ('w:color', color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def section(doc, title):
    p = para(doc, before=8, after=3, keep_next=True)
    style_run(p.add_run(title), size=11.5, bold=True, color=ACCENT, caps=True)
    bottom_border(p)


def bullet(doc, text, lead=None):
    p = para(doc, after=1.5, line=1.03)
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Inches(0.18), Inches(-0.14)
    style_run(p.add_run('•  '), color=ACCENT)
    if lead:
        style_run(p.add_run(lead + ': '), bold=True)
    style_run(p.add_run(text))


def build(out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.65)
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = FONT, Pt(BODY_PT)

    style_run(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(NAME), size=22, bold=True, color=ACCENT)
    style_run(para(doc, before=2, align=WD_ALIGN_PARAGRAPH.CENTER).add_run(HEADLINE), size=11.5, bold=True)
    p = para(doc, before=2, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    style_run(p.add_run(CONTACT), size=10, color=MUTED)
    bottom_border(p, size=12)

    section(doc, 'Executive Summary')
    for i, s in enumerate(SUMMARY):
        style_run(para(doc, after=3 if i == 0 else 1, line=1.05).add_run(s))

    section(doc, 'Selected Leadership Results')
    for lead, text in RESULTS:
        bullet(doc, text, lead)

    section(doc, 'Core Expertise')
    style_run(para(doc, after=1, line=1.1).add_run('  |  '.join(EXPERTISE)))

    section(doc, 'Professional Experience')
    for org, loc, title, dates, summary, bullets in JOBS:
        p = para(doc, before=6, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(org), bold=True, color=ACCENT)
        style_run(p.add_run(f'  |  {loc}'), color=MUTED)
        style_run(p.add_run('\t' + dates), bold=True)
        style_run(para(doc, after=1, keep_next=True).add_run(title), bold=True, italic=True)
        if summary:
            style_run(para(doc, after=1.5, keep_next=True).add_run(summary), size=10, italic=True, color=MUTED)
        for b in bullets:
            bullet(doc, b)

    section(doc, 'Technology & Program Innovation')
    for lead, text in INNOVATION:
        bullet(doc, text, lead)

    section(doc, 'Education')
    for degree, school, year in EDUCATION:
        p = para(doc, after=1)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        style_run(p.add_run(degree), bold=True)
        style_run(p.add_run(f'  |  {school}'))
        if year:
            style_run(p.add_run('\t' + year))

    doc.core_properties.title = 'Clavel Lazarre, MBA — Executive Resume'
    doc.core_properties.author = 'Clavel Lazarre'
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build(sys.argv[1])
